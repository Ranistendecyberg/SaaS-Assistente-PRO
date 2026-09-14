from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Manual_do_Usuario_SaaS_Assistente_PRO_2.1.4.docx"
ICON = ROOT / "src" / "assets" / "icon.png"

BLUE = "2563EB"
DARK = "0F172A"
MID = "475569"
LIGHT = "EFF6FF"
GRID = "D9D9D9"
ORANGE = "EA580C"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def add_steps(doc, items):
    for number, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.add_run(f"{number}.  ")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def add_note(doc, title, text, color=ORANGE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    p.add_run(title + " ").bold = True
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = header
        shade(cell, DARK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
        set_cell_border(cell)
        cell_margin(cell)
        if widths:
            cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if ridx % 2:
                shade(cells[i], "F8FAFC")
            set_cell_border(cells[i])
            cell_margin(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def new_page(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string("000000")
    style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(5)
styles["List Bullet"].font.size = Pt(10)
styles["List Number"].font.size = Pt(10)

header = section.header.paragraphs[0]
header.text = "SaaS Assistente PRO  |  Manual do Usuário  |  versão 2.1.4"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
add_page_number(section.footer.paragraphs[0])

# Capa
doc.add_paragraph().paragraph_format.space_after = Pt(28)
if ICON.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ICON), width=Inches(1.05))
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Manual do Usuário")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SaaS Assistente PRO")
r.bold = True
r.font.size = Pt(23)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Pesquisas SSI e TSI, WhatsApp, dashboards, licenças e conta empresarial")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string(MID)
doc.add_paragraph().paragraph_format.space_after = Pt(70)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Versão documentada 2.1.4")
r.bold = True
r.font.size = Pt(12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Uma solução do portfólio Grupo HAGO").italic = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Setembro de 2026")

new_page(doc)
add_heading(doc, "Como usar este manual", 1)
add_para(doc, "Este documento explica o funcionamento operacional do SaaS Assistente PRO, os botões da versão 2.1.4 e as regras que afetam licença, sincronização, disparos, relatórios, conta empresarial e cobrança. Os nomes dos botões aparecem exatamente como no programa sempre que possível.")
add_note(doc, "Regra confirmada", "Nos dias 1, 2 e 3 de cada mês, os auditores SSI e TSI consultam o mês atual e o mês anterior. A partir do dia 4, consultam somente o mês atual.", BLUE)
add_heading(doc, "Sumário de assuntos", 2)
topics = [
    "1. Visão geral e requisitos", "2. Primeiro acesso e abertura do sistema",
    "3. Licenças, validade, limites e cobrança", "4. Navegação e indicadores laterais",
    "5. Motor WhatsApp", "6. Gera Lista Reenvio e myHonda", "7. Auditores SSI e TSI",
    "8. Fila, busca, seleção e disparo", "9. Modelos de mensagem",
    "10. Relatórios e dashboards", "11. Configuração de lojas e consultores",
    "12. Conta Empresarial e computadores", "13. Cobrança por PIX e boleto",
    "14. Chaves de acesso", "15. Tutorial, sugestões, termos e atualizações",
    "16. Regras de segurança e privacidade", "17. Solução de problemas",
    "18. Rotina recomendada e checklist", "19. Glossário", "20. Suporte"
]
add_bullets(doc, topics)

new_page(doc)
add_heading(doc, "1 Visão geral e requisitos", 1)
add_para(doc, "O SaaS Assistente PRO apoia concessionárias na recuperação e no acompanhamento de pesquisas de satisfação da Honda. Ele reúne dados do myHonda, identifica clientes elegíveis, auxilia o envio de links oficiais pelo WhatsApp e consolida resultados SSI e TSI em dashboards gerenciais.")
add_heading(doc, "O que o sistema faz", 2)
add_bullets(doc, [
    "Mantém uma sessão separada do myHonda para consultar filas e relatórios de pesquisas.",
    "Mantém uma sessão separada do WhatsApp Web para envio individual ou sequencial.",
    "Executa os auditores SSI e TSI e guarda localmente o histórico necessário aos painéis.",
    "Reduz reenvios indevidos ao considerar histórico de envio e pesquisas já respondidas.",
    "Apresenta painéis TSI, SSI e um relatório gerencial geral, com filtros e exportação em PDF.",
    "Controla licença, limite diário, bônus, computadores vinculados e cobrança empresarial."
])
add_heading(doc, "Requisitos", 2)
add_bullets(doc, [
    "Windows 10 ou Windows 11, com internet estável.",
    "Acesso válido ao myHonda para a operação das pesquisas.",
    "Um WhatsApp de atendimento autorizado, preferencialmente o número oficial da concessionária.",
    "Conta/licença ativa no SaaS Assistente PRO.",
    "Permissão interna e base legal para contatar os clientes selecionados."
])
add_note(doc, "Importante", "O sistema depende de serviços externos, principalmente myHonda/Salesforce, WhatsApp Web, Supabase e Mercado Pago. Lentidão, indisponibilidade, mudança de tela ou sessão expirada nesses serviços pode afetar temporariamente a operação.")

new_page(doc)
add_heading(doc, "2 Primeiro acesso e abertura do sistema", 1)
add_heading(doc, "Sequência de abertura", 2)
add_steps(doc, [
    ("Cópia de segurança local. ", "O programa inicia uma cópia de segurança silenciosa dos dados locais."),
    ("Termos de uso. ", "No primeiro uso, é necessário ler e aceitar os termos. Recusar encerra o programa."),
    ("Primeiro acesso. ", "Escolha Nova conta, Computador adicional ou Já tenho conta — Entrar."),
    ("Validação da licença. ", "O servidor confere a instalação, a validade e os limites. Falhas momentâneas de rede podem gerar nova tentativa."),
    ("Avisos de serviço. ", "Um comunicado importante ou modo de manutenção pode ser exibido antes da tela principal."),
    ("Tela principal. ", "O menu lateral e os indicadores mostram o plano, os dias restantes e o consumo diário."),
    ("Atualização. ", "O aplicativo verifica se existe uma versão mais recente publicada.")
])
add_heading(doc, "Botões do primeiro acesso", 2)
add_table(doc, ["Botão", "Função"], [
    ["Nova conta", "Cadastra empresa ou pessoa, proprietário, telefone, CPF/CNPJ, e-mail e senha. O e-mail precisa ser confirmado antes da criação do teste."],
    ["Computador adicional", "Ativa esta máquina com um código temporário emitido na Conta Empresarial de outro computador autorizado."],
    ["Já tenho conta — Entrar", "Abre o acesso por e-mail e senha para recuperar a conta empresarial existente."],
    ["Criar conta e enviar código", "Cria o usuário e envia a confirmação ao e-mail informado."],
    ["Confirmar e iniciar teste", "Confere o código recebido e, se válido, cria a conta empresarial e o teste gratuito."],
    ["Já confirmei pelo link do e-mail", "Verifica a confirmação feita no link recebido e conclui a entrada."],
    ["Ativar este computador", "Valida o código de computador adicional. O código expira em 15 minutos e só pode ser usado uma vez."],
    ["Sair", "Fecha o assistente de primeiro acesso."]
], [2.05, 5.0])
add_heading(doc, "Recuperação de acesso", 2)
add_para(doc, "Na tela de login, Esqueci minha senha abre a recuperação. Informe o e-mail, use Enviar código de recuperação, digite o código recebido, a nova senha e a confirmação, e finalize em Validar código e alterar senha. A senha deve ter pelo menos oito caracteres. Se a conta exigir autenticação em duas etapas, também será pedido o código atual de seis números do aplicativo autenticador.")
add_heading(doc, "Confirmações de segurança", 2)
add_para(doc, "Algumas ações protegidas podem exigir uma confirmação adicional. Na janela Confirmação de segurança, informe o código atual do aplicativo autenticador e use Confirmar código; Cancelar interrompe a operação. Na confirmação por e-mail, Reenviar código solicita outro código, Confirmar operação valida o código informado e Cancelar abandona a ação. Um Comunicado do Sistema é fechado por Entendido e Ciente depois da leitura.")
add_heading(doc, "3 Licenças, validade, limites e cobrança", 1)
add_heading(doc, "Tipos de licença", 2)
add_table(doc, ["Situação", "Como funciona", "Limite padrão"], [
    ["Teste grátis", "Válido por 2 dias após a confirmação do e-mail e criação da conta. Permite avaliar as funções principais.", "6 mensagens por dia"],
    ["Plano PRO ativo", "Assinatura mensal vinculada à conta empresarial e aos computadores faturáveis.", "30 mensagens por dia, salvo configuração administrativa"],
    ["Bônus de mensagens", "Quantidade extra concedida à licença. Enquanto houver saldo, o indicador lateral mostra o bônus disponível.", "Conforme o saldo concedido"],
    ["Chave de acesso", "Chave promocional, cortesia ou regularização que pode acrescentar validade ou liberar o acesso conforme sua configuração.", "Conforme a chave"],
    ["Vencida", "Bloqueia os envios e solicita renovação, PIX ou chave válida.", "Envios bloqueados"]
], [1.35, 4.1, 1.55])
add_heading(doc, "Como a validade é exibida", 2)
add_bullets(doc, [
    "Teste Grátis: X dias aparece para a fase de avaliação.",
    "Plano PRO: X dias aparece para uma assinatura ativa.",
    "Quando resta um dia, o programa informa Último dia e indica que a validade termina às 22h.",
    "Plano: Vencido indica que a licença precisa ser renovada ou regularizada.",
    "Envios Hoje mostra enviados/limite; Bônus de Licença aparece quando existem mensagens extras."
])
add_heading(doc, "Como são cobradas as licenças", 2)
add_para(doc, "A cobrança empresarial é consolidada. O valor considera o preço-base da conta e os computadores adicionais faturáveis. A referência comercial padrão do sistema é R$ 300,00 de base mais R$ 50,00 por computador adicional, porém preços, descontos, isenções e condições podem ser configurados administrativamente; por isso, o valor válido é sempre o confirmado pelo servidor na tela de cobrança.")
add_note(doc, "Atenção ao valor", "O valor mostrado antes de gerar a cobrança é uma estimativa. A fatura definitiva é recalculada e confirmada pelo servidor no momento da emissão.")
add_heading(doc, "Regras de limite", 2)
add_bullets(doc, [
    "O limite diário pertence à licença e é validado no servidor antes do envio.",
    "O administrador pode configurar limites diferentes do padrão.",
    "Existe também um limite por lote, configurável entre 1 e 1.000 mensagens.",
    "O sistema reserva a cota antes do clique de envio e confirma o consumo quando conclui.",
    "Em uma falha comprovadamente anterior ao envio, a reserva pode ser liberada. Se houver dúvida após o clique no WhatsApp, a mensagem permanece contabilizada para evitar duplicidade."
])

new_page(doc)
add_heading(doc, "4 Navegação e indicadores laterais", 1)
add_table(doc, ["Menu", "O que abre"], [
    ["Gera Lista Reenvio", "Área principal do myHonda, auditores, filas SSI/TSI, modelos e disparos."],
    ["Motor WhatsApp", "WhatsApp Web usado pelo programa. Permite ler o QR Code e recarregar a sessão."],
    ["Relatórios / Dashboard", "Abas de resultados TSI, SSI e Relatório Gerencial Geral."],
    ["Configuração de Lojas", "Cadastro de lojas, metas TSI, consultores e limpeza do histórico de envios."],
    ["Conta Empresarial", "Empresa, unidades, computadores, perfil de acesso, estimativa e cobrança."],
    ["Validar Chave de Acesso", "Tela de renovação, geração de PIX e validação de chave."],
    ["Tutorial do Sistema", "Resumo dos primeiros passos dentro do próprio programa."],
    ["Sugestões de Melhoria", "Formulário para enviar sugestões à equipe do produto."],
    ["Sobre o Sistema", "Versão instalada, atualização automática e termos de uso."]
], [2.2, 4.85])
add_para(doc, "O rodapé do menu mostra o plano e a quantidade de dias, além do uso diário ou do bônus. Esses números vêm da validação da licença; após uma renovação, use Atualizar dados ou reinicie o programa se a informação ainda não tiver sido renovada na tela.")

new_page(doc)
add_heading(doc, "5 Motor WhatsApp", 1)
add_para(doc, "O Motor WhatsApp é o navegador exclusivo do WhatsApp Web. Ele não compartilha a sessão do myHonda. A conexão fica armazenada no perfil local do aplicativo para reduzir a necessidade de ler o QR Code a cada abertura.")
add_heading(doc, "Primeira conexão", 2)
add_steps(doc, [
    "Abra Motor WhatsApp no menu lateral.",
    "No celular, abra o WhatsApp oficial da concessionária e acesse Aparelhos conectados.",
    "Leia o QR Code exibido pelo programa.",
    "Aguarde a lista de conversas aparecer antes de voltar à fila de pesquisas."
])
add_table(doc, ["Botão ou estado", "Função"], [
    ["Recarregar WhatsApp", "Reabre o WhatsApp Web e tenta recuperar uma sessão travada, desconectada ou incompleta."],
    ["Envio disponível na Lista de Clientes", "Indicador: os envios operacionais devem ser iniciados em Gera Lista Reenvio, e não nesta tela."],
    ["Link Enviado", "Confirma que a automação encontrou a mensagem e acionou o envio."],
    ["Número Inválido", "O WhatsApp informou que o telefone não pode receber a mensagem."],
    ["Falha ao Enviar", "A página não ficou pronta, a sessão caiu ou o elemento esperado não foi encontrado."]
], [2.5, 4.55])
add_note(doc, "Uso responsável", "Utilize somente contatos autorizados e respeite as políticas do WhatsApp e a legislação de proteção de dados. Volumes ou padrões considerados abusivos podem causar restrições aplicadas pelo próprio WhatsApp.")

new_page(doc)
add_heading(doc, "6 Gera Lista Reenvio e myHonda", 1)
add_para(doc, "Esta é a área central de trabalho. Nela, o usuário entra no myHonda, atualiza as bases, acompanha os auditores, escolhe SSI ou TSI, localiza clientes e envia a pesquisa apropriada.")
add_heading(doc, "Botões de sincronização", 2)
add_table(doc, ["Botão", "O que faz", "Quando usar"], [
    ["Atualizar myHonda", "Força nova sincronização das filas e dos relatórios mensais SSI e TSI. Os indicadores mostram o andamento.", "No início da rotina e quando houver alteração recente no myHonda."],
    ["Histórico Anual", "Aplica o período anual nos relatórios e importa o histórico disponível do ano corrente.", "Na implantação, para recompor histórico ou quando os painéis anuais precisarem ser atualizados."],
    ["Conversar com o Destacado", "Abre a conversa do cliente destacado no WhatsApp. O usuário revisa a conversa; a mensagem não é enviada automaticamente por esse botão.", "Para atendimento individual e conferência."],
    ["Enviar Pesquisa Selecionada", "Inicia o envio da seleção válida, após conferir licença, cota, conexão e dados necessários.", "Depois de revisar pesquisa, cliente e modelo de mensagem."]
], [1.75, 3.4, 1.9])
add_heading(doc, "Login no myHonda", 2)
add_steps(doc, [
    "Abra Gera Lista Reenvio.",
    "Entre no myHonda dentro do painel do sistema.",
    "Aguarde o indicador myHonda: Autenticado e a preparação dos relatórios.",
    "Acompanhe Auditor TSI e Auditor SSI até que as bases estejam prontas.",
    "Quando a sincronização terminar, a fila é gerada automaticamente."
])
add_note(doc, "Sessão expirada", "Se o myHonda voltar à tela de login, faça a autenticação novamente. Se o portal ficar incompleto, use Atualizar myHonda depois que a página estiver autenticada.")

new_page(doc)
add_heading(doc, "7 Auditores SSI e TSI", 1)
add_para(doc, "Os auditores são rotinas de leitura dos relatórios do myHonda/Salesforce. Eles não atribuem notas e não alteram respostas: coletam os dados disponíveis, normalizam campos, eliminam duplicidades e atualizam o histórico local usado na fila e nos dashboards.")
add_table(doc, ["Auditor", "Área", "Principais usos"], [
    ["SSI", "Vendas e entrega", "Identificar pesquisas de satisfação de vendas, consolidar respostas, vendedores, lojas, modalidade/modelo quando disponível e alimentar o painel SSI."],
    ["TSI", "Pós-vendas e serviço", "Identificar pesquisas técnicas/serviço, cruzar ordens de serviço, respostas, consultores e categorias, impedir contato indevido com respondidos e alimentar o painel TSI."]
], [1.0, 2.0, 4.05])
add_heading(doc, "Regra dos três primeiros dias", 2)
add_para(doc, "Nos dias 1, 2 e 3 do mês, tanto o Auditor SSI quanto o Auditor TSI consultam o mês atual e o mês anterior. Essa janela evita que respostas registradas no fim do mês anterior fiquem ausentes quando o usuário não sincronizou o programa nos últimos dias daquele mês. A partir do dia 4, a coleta mensal volta a considerar somente o mês atual.")
add_note(doc, "Por que pode demorar", "Nos três primeiros dias há mais informações para o myHonda/Salesforce preparar e entregar. Enquanto o portal processa o período ampliado, as listas e os dashboards podem levar mais tempo para atualizar. Isso, por si só, não representa defeito do SaaS Assistente PRO.", BLUE)
add_heading(doc, "Leitura dos indicadores", 2)
add_bullets(doc, [
    "Buscando ou Sincronizando: o portal ainda está sendo consultado.",
    "Filtro aplicado, aguardando conclusão: o período foi solicitado e o relatório ainda está processando.",
    "Relatório concluído, importando dados: a leitura terminou e os registros estão sendo gravados/mesclados.",
    "X Respostas: quantidade reconhecida para o período atual.",
    "Instabilidade detectada, recuperando: o auditor encontrou falha temporária e fará nova tentativa.",
    "Timeout: o myHonda/Salesforce não concluiu no tempo esperado; confira a sessão e tente atualizar novamente."
])
add_heading(doc, "Histórico anual", 2)
add_para(doc, "A coleta anual percorre um volume muito maior e acompanha o relatório até que ele se estabilize. Em contas com muitos registros, pode levar vários minutos. Evite fechar o programa, trocar filtros no myHonda ou interromper a internet durante esse processo.")

new_page(doc)
add_heading(doc, "8 Fila, busca, seleção e disparo", 1)
add_heading(doc, "Como interpretar a fila", 2)
add_bullets(doc, [
    "Fila de Disparo mostra a quantidade de registros visíveis após os filtros e a busca.",
    "A caixa de busca localiza clientes pelo nome.",
    "O tipo de pesquisa separa SSI de Vendas e TSI de Pós-Vendas.",
    "Marcar a caixa do registro seleciona a pesquisa para disparo; destacar a linha prepara a conversa individual.",
    "Registros já enviados podem aparecer bloqueados. Pesquisas TSI já respondidas também são protegidas contra novo disparo.",
    "A fila depende dos dados recebidos do myHonda e das configurações locais de lojas e consultores."
])
add_heading(doc, "Procedimento seguro de envio", 2)
add_steps(doc, [
    "Confirme que myHonda e WhatsApp estão autenticados.",
    "Escolha SSI ou TSI e aguarde a lista correta.",
    "Busque o cliente e confira se a pesquisa é a autorizada.",
    "Escolha ou revise o modelo de mensagem.",
    "Marque os registros desejados, respeitando a autorização e o limite exibido.",
    "Clique em Enviar Pesquisa Selecionada.",
    "Acompanhe cada cliente no status e, ao final, revise o Resumo de Envios Não Concluídos."
])
add_heading(doc, "Regras durante o disparo", 2)
add_bullets(doc, [
    "Antes de começar, o sistema valida status da licença, limite diário, bônus, limite do lote e conexão do WhatsApp.",
    "Os envios são processados um a um, com intervalo aleatório aproximado de 6,5 a 11,5 segundos entre sucessos.",
    "No SSI, quando o contato não está na fila, o programa pode abrir a ficha do cliente no myHonda para obter telefone e outros dados necessários.",
    "O WhatsApp pode levar até cerca de 40 segundos para preparar a conversa e localizar a mensagem/botão de envio.",
    "Telefone ausente, ficha indisponível, número inválido, página não carregada ou sessão desconectada impedem a conclusão daquele cliente.",
    "Ao interromper o processo, mensagens com resultado ambíguo podem permanecer contadas por segurança."
])
add_heading(doc, "Resumo de disparo", 2)
add_para(doc, "Quando existem clientes não notificados, a janela Resumo de Envios Não Concluídos lista as inconsistências. Copiar Lista envia o conteúdo para a área de transferência; Entendido fecha a janela. Corrija os cadastros no myHonda antes de tentar novamente.")

new_page(doc)
add_heading(doc, "9 Modelos de mensagem", 1)
add_para(doc, "O editor permite manter textos reutilizáveis. Antes de cada disparo, confira saudação, contexto, link e identificação da concessionária.")
add_table(doc, ["Controle", "Função"], [
    ["Lista de modelos", "Carrega um texto já salvo para leitura ou edição."],
    ["Salvar Novo", "Cria um novo modelo com o texto atual. O sistema solicita um nome para identificá-lo."],
    ["Salvar", "Atualiza o modelo atualmente selecionado."],
    ["Excluir", "Remove o modelo selecionado após a confirmação."],
    ["+ [NOME]", "Insere a variável que será substituída pelo nome do cliente."],
    ["+ [LINK]", "Insere a variável do link oficial da pesquisa, gerado para o registro selecionado."],
    ["Editor de mensagem", "Área onde o texto é escrito e revisado antes do envio."]
], [1.75, 5.3])
add_note(doc, "Antes de salvar", "Mantenha [NOME] e [LINK] exatamente entre colchetes. Não cole links genéricos no lugar de [LINK] se a intenção for usar o endereço específico da pesquisa.")

new_page(doc)
add_heading(doc, "10 Relatórios e dashboards", 1)
add_para(doc, "Relatórios / Dashboard possui três abas. Os dados vêm do histórico local produzido pelos auditores; por isso, sincronize o myHonda antes de interpretar resultados recentes.")
add_heading(doc, "Painel de Resultados TSI", 2)
add_table(doc, ["Controle", "Função"], [
    ["Concessionária", "Filtra uma unidade cadastrada ou todas."],
    ["Consultor Técnico", "Filtra o profissional de pós-vendas."],
    ["Mês", "Seleciona o mês; por padrão, procura o mês atual ou o mais recente disponível."],
    ["Categoria Produto", "Recorta os resultados pela categoria disponível no relatório."],
    ["Atualizar Painel", "Recarrega a base local e reaplica os filtros."],
    ["PDF Diretoria", "Gera um relatório executivo em PDF com a visão filtrada."]
], [1.75, 5.3])
add_para(doc, "O TSI apresenta volume de respostas, indicadores de satisfação, evolução, desempenho por consultor/loja e métricas válidas para as notas encontradas. Valores ausentes ou inválidos ficam fora do denominador das notas, em vez de serem classificados automaticamente como detratores.")
add_heading(doc, "Painel de Resultados SSI", 2)
add_table(doc, ["Controle", "Função"], [
    ["Concessionária", "Filtra uma unidade de vendas."],
    ["Mês", "Seleciona um mês específico ou todos."],
    ["Modalidade", "Filtra a modalidade de compra disponível no relatório."],
    ["Vendedor", "Filtra o vendedor; a lista acompanha loja e mês escolhidos."],
    ["Atualizar Painel", "Lê novamente a base SSI local."],
    ["PDF Diretoria", "Exporta a visão executiva SSI em PDF."]
], [1.75, 5.3])
add_para(doc, "O SSI consolida respostas de vendas e indicadores como satisfação, recomendação, volume, evolução e desempenho por vendedor, conforme os campos disponibilizados pelo myHonda.")
add_heading(doc, "Relatório Gerencial Geral", 2)
add_para(doc, "Permite alternar Departamento entre Pós-Vendas TSI e Vendas SSI, filtrar Concessionária, Mês e Consultor/Vendedor, usar Atualizar e gerar o Relatório Executivo PDF. A visão compara a conversão dos envios pelo WhatsApp com os canais tradicionais disponíveis nos dados.")
add_heading(doc, "Interpretação de recomendação", 2)
add_table(doc, ["Nota válida", "Classificação"], [["9 ou 10", "Promotor"], ["7 ou 8", "Neutro"], ["0 a 6", "Detrator"], ["Ausente, inválida ou fora de 0 a 10", "Sem nota, fora do cálculo de classificação"]], [2.8, 4.25])
add_note(doc, "Painel desatualizado", "Se uma resposta acabou de surgir no myHonda, primeiro execute Atualizar myHonda e aguarde os dois auditores terminarem. Só depois use Atualizar Painel. Nos três primeiros dias, essa sequência pode demorar mais por causa da janela de dois meses.")

new_page(doc)
add_heading(doc, "11 Configuração de lojas e consultores", 1)
add_heading(doc, "Lojas", 2)
add_table(doc, ["Campo ou botão", "Função"], [
    ["Nome da loja", "Nome amigável exibido nos filtros e relatórios."],
    ["CNPJ ou código", "Identificador usado para relacionar registros do myHonda à loja. Digite somente números quando solicitado."],
    ["Meta TSI", "Meta de referência da unidade para análises TSI."],
    ["Adicionar Loja", "Grava uma nova loja após validar os campos."],
    ["Salvar Alterações", "Atualiza a loja selecionada."],
    ["Remover Selecionada", "Exclui a configuração escolhida após confirmação."]
], [2.05, 5.0])
add_heading(doc, "Consultores", 2)
add_table(doc, ["Campo ou botão", "Função"], [
    ["Nome do consultor", "Nome exibido nos filtros e relatórios."],
    ["CPF", "Identificador de vínculo com os registros; digite somente números."],
    ["Adicionar Consultor", "Inclui um profissional."],
    ["Salvar Alterações", "Atualiza o profissional selecionado."],
    ["Remover Selecionado", "Exclui o cadastro local após confirmação."]
], [2.05, 5.0])
add_heading(doc, "Limpar Histórico de Envios", 2)
add_para(doc, "Apaga a marcação local usada para identificar pesquisas já enviadas. Use apenas quando houver motivo operacional confirmado, pois registros antigos podem voltar a aparecer como elegíveis e aumentar o risco de reenvio. O histórico de respostas dos auditores é uma base diferente.")
add_note(doc, "Cuidado", "Limpar o histórico não cancela mensagens no WhatsApp e não apaga respostas no myHonda. Antes de usar, confirme internamente quais registros poderão ser reenviados.")

new_page(doc)
add_heading(doc, "12 Conta Empresarial e computadores", 1)
add_para(doc, "A Conta Empresarial concentra empresa, unidades, assinatura, computadores e cobrança. Ao abrir, o sistema pode solicitar e-mail e senha. Atualizar dados força uma nova leitura do servidor.")
add_heading(doc, "Perfis e permissões", 2)
add_table(doc, ["Perfil", "Pode visualizar", "Pode administrar", "Pode tornar principal"], [
    ["Proprietário", "Tudo", "Unidades, códigos, dispositivos e cobrança", "Sim"],
    ["Administrador", "Tudo", "Códigos e dispositivos, conforme autorização", "Não"],
    ["Operador", "Resumo da conta", "Não; ações ficam desabilitadas", "Não"]
], [1.4, 1.65, 2.8, 1.2])
add_heading(doc, "Áreas da tela", 2)
add_bullets(doc, [
    "Resumo: situação da assinatura, vencimento, computadores faturáveis e valor estimado.",
    "Unidades e documentos: matriz/filiais, tipo, CPF/CNPJ e status.",
    "Adicionar computador: escolhe a unidade e gera um código temporário.",
    "Computadores vinculados: mostra computador, unidade, classe, situação, versão e último acesso.",
    "Cobrança consolidada: abre os dados fiscais e as formas de pagamento."
])
add_heading(doc, "Botões", 2)
add_table(doc, ["Botão", "Regra"], [
    ["Atualizar dados", "Busca no servidor as informações mais recentes da conta."],
    ["Gerar código de ativação", "Emite código no formato PC-XXXX-XXXX para a unidade escolhida. Válido por 15 minutos e uso único. Requer perfil autorizado e assinatura que permita vínculo."],
    ["Copiar código", "Copia o código emitido. O código é mostrado somente uma vez na tela."],
    ["Tornar Principal", "Transfere a classe de computador principal para o adicional selecionado. Somente o proprietário pode executar."],
    ["Programar remoção", "Agenda a retirada do computador adicional selecionado conforme as regras da conta."],
    ["Cancelar remoção", "Desfaz uma remoção programada ainda não efetivada."],
    ["Abrir cobrança e pagamentos", "Abre a fatura, perfil fiscal, PIX, boleto e consulta de pagamento."]
], [2.15, 4.9])
add_note(doc, "Computador adicional", "O código liga uma instalação específica à conta e à unidade escolhida. Compartilhe-o apenas com o responsável por aquela máquina. A inclusão pode alterar o valor faturável mostrado na próxima cobrança.")

new_page(doc)
add_heading(doc, "13 Cobrança por PIX e boleto", 1)
add_para(doc, "Somente o proprietário pode gerar pagamentos ou consultar diretamente o provedor. Administradores e operadores podem receber visualização somente leitura conforme as permissões da conta.")
add_heading(doc, "Dados para emissão", 2)
add_para(doc, "Preencha os dados fiscais e de contato solicitados e use Salvar dados de cobrança. O servidor usa essas informações na emissão do PIX ou boleto. Revise razão social/nome, documento, e-mail, telefone e endereço antes de gerar.")
add_table(doc, ["Botão", "Função e regra"], [
    ["Salvar dados de cobrança", "Valida e grava o perfil de faturamento no servidor."],
    ["Gerar PIX", "Confere novamente fatura e computadores e solicita um PIX ao Mercado Pago."],
    ["Gerar boleto", "Confere os mesmos dados e solicita boleto quando essa modalidade estiver disponível."],
    ["Copiar código", "Copia o PIX Copia e Cola ou código disponível."],
    ["Abrir pagamento", "Abre o endereço seguro retornado pelo provedor."],
    ["Já paguei — consultar Mercado Pago", "Consulta manualmente o provedor. Há intervalo mínimo aproximado de 30 segundos entre consultas."],
    ["Fechar", "Fecha a janela sem cancelar cobranças já emitidas."]
], [2.4, 4.65])
add_heading(doc, "Confirmação", 2)
add_bullets(doc, [
    "Gerar um código não confirma pagamento. A licença é renovada somente quando o servidor reconhece a aprovação.",
    "O resumo da conta pode se atualizar automaticamente em ciclos aproximados de 15 segundos, sem consultar continuamente o provedor.",
    "Depois da aprovação, use Atualizar dados na Conta Empresarial para conferir a nova vigência.",
    "Se aparecer Pagamento em conferência, não pague novamente. Aguarde e contate o suporte.",
    "Nunca envie senhas bancárias ou códigos de autenticação ao suporte. O programa não solicita credenciais do seu banco."
])

new_page(doc)
add_heading(doc, "14 Chaves de acesso", 1)
add_para(doc, "Validar Chave de Acesso atende chaves promocionais, cortesia ou regularização emitidas pela administração. A tela também oferece o fluxo de PIX para uma licença vencida.")
add_table(doc, ["Botão", "Função"], [
    ["Consultar valor e gerar PIX", "Solicita ao servidor o valor atualizado e cria o pagamento disponível."],
    ["Tenho uma Chave de Ativação", "Mostra o campo para colar uma chave como PRO-12345."],
    ["Validar Chave", "Envia a chave ao servidor. Se válida e aplicável à instalação, libera o benefício configurado."],
    ["Copiar PIX Copia e Cola", "Copia o código PIX gerado."],
    ["Voltar", "Retorna à opção anterior sem ativar ou cancelar cobranças."]
], [2.55, 4.5])
add_para(doc, "Uma chave inválida, expirada, já utilizada ou incompatível não será aplicada. Não publique chaves em grupos ou canais abertos, pois podem ser de uso único ou vinculadas a uma instalação.")

new_page(doc)
add_heading(doc, "15 Tutorial, sugestões, termos e atualizações", 1)
add_heading(doc, "Tutorial do Sistema", 2)
add_para(doc, "Apresenta um guia resumido sobre conexão do WhatsApp, geração de listas, modelos de mensagem, dashboards, configurações e cobrança. O presente manual é a referência mais completa.")
add_heading(doc, "Sugestões de Melhoria", 2)
add_para(doc, "Digite uma ideia no campo Escreva sua sugestão aqui e clique em Enviar Sugestão. Durante o envio, o botão mostra Enviando. Não inclua senhas, códigos bancários, dados excessivos de clientes ou informações sigilosas.")
add_heading(doc, "Termos de uso", 2)
add_table(doc, ["Controle", "Função"], [
    ["Declaro que li...", "Marca a concordância necessária no primeiro acesso."],
    ["Concordar e Continuar", "Registra o aceite e segue para o sistema."],
    ["Recusar e Sair", "Não registra aceite e fecha o programa."],
    ["Fechar", "Fecha a consulta dos termos quando aberta pela tela Sobre."]
], [2.45, 4.6])
add_heading(doc, "Sobre o Sistema e atualização", 2)
add_table(doc, ["Elemento", "Função"], [
    ["Versão Instalada", "Mostra a versão em uso. Este manual corresponde à 2.1.4."],
    ["Baixar e Atualizar Automaticamente", "Baixa o instalador publicado, valida sua integridade e inicia a atualização segura."],
    ["Termos de Uso e Licença", "Abre o texto vigente para consulta."],
    ["Status", "Informa se existe versão nova, se o sistema está atualizado ou se não foi possível consultar."]
], [2.6, 4.45])
add_note(doc, "Windows SmartScreen", "Como o instalador pode não possuir assinatura digital de editor, o Windows pode exibir um aviso. Confirme que o arquivo veio do site/repositório oficial e que a versão e a validação do atualizador estão corretas antes de prosseguir.")

new_page(doc)
add_heading(doc, "16 Regras de segurança e privacidade", 1)
add_bullets(doc, [
    "Use contas individuais e não compartilhe senhas do myHonda, e-mail ou Conta Empresarial.",
    "Proteja os códigos de computador adicional e chaves de acesso; ambos podem conceder benefícios ou vínculo à conta.",
    "Mantenha Windows e SaaS Assistente PRO atualizados.",
    "Selecione somente clientes autorizados e envie apenas a pesquisa relacionada ao atendimento correto.",
    "Não exporte relatórios para pastas públicas sem necessidade. PDFs podem conter dados comerciais e indicadores internos.",
    "A sessão do myHonda e a sessão do WhatsApp ficam em perfis locais separados. Evite usar computadores compartilhados sem controle de acesso.",
    "O histórico local serve à proteção contra duplicidade e aos painéis. Faça cópias de segurança antes de manutenção ou troca de computador.",
    "O pagamento é confirmado pelo servidor/provedor. O suporte nunca precisa de sua senha bancária."
])
add_heading(doc, "Responsabilidade operacional", 2)
add_para(doc, "A automação reduz trabalho repetitivo, mas a decisão de contatar o cliente continua sendo da concessionária. O usuário deve revisar destinatário, pesquisa, mensagem e permissão antes de iniciar o envio.")

new_page(doc)
add_heading(doc, "17 Solução de problemas", 1)
add_table(doc, ["Sintoma", "Causa provável", "O que fazer"], [
    ["myHonda pede login", "Sessão expirada", "Autentique novamente e depois use Atualizar myHonda."],
    ["Atualização demora nos dias 1 a 3", "Consulta do mês atual e anterior", "Aguarde o myHonda/Salesforce concluir; não é necessariamente falha."],
    ["Auditor mostra timeout", "Portal lento, relatório não estabilizou ou sessão incompleta", "Confira login e internet; tente Atualizar myHonda novamente."],
    ["Dashboard não mostra resposta recente", "Base local ainda não foi sincronizada", "Conclua Atualizar myHonda e depois clique em Atualizar Painel."],
    ["Lista vazia", "Sem elegíveis, filtros ativos, sincronização incompleta ou dados do portal ausentes", "Limpe a busca, confira SSI/TSI, aguarde auditores e revise o myHonda."],
    ["WhatsApp mostra QR Code", "Sessão não conectada", "Leia o QR Code com o número oficial."],
    ["Número inválido", "Telefone ausente/incorreto ou rejeitado pelo WhatsApp", "Corrija o cadastro no myHonda e tente apenas após validar o contato."],
    ["Celular não encontrado na ficha", "Ficha SSI sem telefone reconhecível", "Atualize o cadastro de origem; o resumo final registrará a falha."],
    ["Limite atingido", "Cota diária/lote consumida e sem bônus", "Aguarde a renovação diária ou solicite ajuste/benefício à administração."],
    ["Licença vencida", "Fim do teste ou período pago", "Abra cobrança, gere PIX/boleto ou valide uma chave."],
    ["Pagamento em conferência", "Confirmação não conclusiva", "Não pague de novo; aguarde e contate o suporte."],
    ["Atualizador falha", "Rede, arquivo incompleto ou validação de segurança", "Tente novamente; se persistir, baixe somente da fonte oficial e contate o suporte."],
    ["Erro No module named PyQt6.sip", "Instalador antigo foi empacotado sem dependência", "Instale a versão 2.1.4 corrigida, publicada no canal oficial."],
    ["Site ou serviço indisponível", "Falha externa ou manutenção", "Aguarde o restabelecimento e tente novamente sem repetir cobranças/envios incertos."]
], [1.65, 2.45, 3.0])

new_page(doc)
add_heading(doc, "18 Rotina recomendada e checklist", 1)
add_heading(doc, "Rotina diária", 2)
add_steps(doc, [
    "Abra o programa e confira plano, dias restantes e cota.",
    "Abra Motor WhatsApp e confirme que a sessão está conectada.",
    "Abra Gera Lista Reenvio e autentique o myHonda, se necessário.",
    "Clique em Atualizar myHonda e espere myHonda, Auditor TSI e Auditor SSI terminarem.",
    "Escolha o tipo correto, revise a fila, busque e selecione os clientes autorizados.",
    "Revise o modelo e faça o disparo. Trate as falhas do resumo final.",
    "Abra os dashboards, clique em Atualizar Painel e analise o período desejado."
])
add_heading(doc, "Cuidados especiais nos dias 1, 2 e 3", 2)
add_bullets(doc, [
    "Reserve mais tempo para a sincronização.",
    "Não interrompa o myHonda enquanto os dois meses são processados.",
    "Espere os auditores concluírem antes de comparar o fechamento do mês anterior.",
    "Só atualize o dashboard depois da mensagem de bases prontas."
])
add_heading(doc, "Antes de fechar", 2)
add_bullets(doc, [
    "Verifique se não há envio em andamento.",
    "Registre e corrija contatos não notificados.",
    "Salve PDFs necessários em local autorizado.",
    "Feche o programa normalmente para preservar os dados locais."
])

new_page(doc)
add_heading(doc, "19 Glossário", 1)
add_table(doc, ["Termo", "Significado neste sistema"], [
    ["SSI", "Pesquisa ligada à experiência de vendas/entrega."],
    ["TSI", "Pesquisa ligada à experiência técnica e de pós-vendas/serviço."],
    ["Auditor", "Rotina que lê e consolida respostas do relatório myHonda/Salesforce."],
    ["Fila de disparo", "Lista de pesquisas potencialmente elegíveis para contato."],
    ["Histórico de envios", "Registro local usado para reduzir reenvios duplicados."],
    ["Histórico anual", "Coleta ampliada do ano corrente para recompor painéis."],
    ["Promotor", "Nota de recomendação válida 9 ou 10."],
    ["Neutro", "Nota de recomendação válida 7 ou 8."],
    ["Detrator", "Nota de recomendação válida de 0 a 6."],
    ["Sem nota", "Resposta sem nota válida de recomendação; não é tratada automaticamente como detrator."],
    ["Top2Box", "Percentual de avaliações nas duas notas superiores da escala considerada pelo indicador."],
    ["Computador principal", "Instalação principal da conta empresarial."],
    ["Computador adicional", "Outra instalação vinculada à conta e possivelmente faturável."],
    ["Reserva de cota", "Bloqueio preventivo de uma vaga de envio antes da tentativa no WhatsApp."],
    ["PIX Copia e Cola", "Código de pagamento retornado pelo provedor."]
], [2.0, 5.05])

new_page(doc)
add_heading(doc, "20 Suporte", 1)
add_para(doc, "Ao pedir ajuda, informe a versão instalada, a tela em que ocorreu o problema, o texto completo do aviso, horário aproximado e, se possível, uma captura de tela sem senhas nem dados bancários.")
add_table(doc, ["Canal", "Endereço"], [
    ["E-mail", "berg.suportetr@gmail.com"],
    ["Site", "https://saas-assistente-pro.berg-suportetr.chatgpt.site"],
    ["Produto", "SaaS Assistente PRO, portfólio Grupo HAGO"]
], [1.5, 5.55])
add_heading(doc, "Informações úteis para o chamado", 2)
add_bullets(doc, [
    "Versão do SaaS Assistente PRO e versão do Windows.",
    "Se o problema ocorreu no myHonda, WhatsApp, licença, cobrança ou dashboard.",
    "Qual botão foi acionado e qual mensagem apareceu.",
    "Se ocorreu nos dias 1 a 3, quando a coleta mensal inclui dois meses.",
    "Se outros sites estavam funcionando e se as sessões estavam autenticadas."
])
add_note(doc, "Nunca envie", "Senha do myHonda, senha do e-mail, senha bancária, código de autenticação ou chave/código de ativação ainda válido.", "DC2626")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(26)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Fim do manual")
r.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("SaaS Assistente PRO 2.1.4  |  Grupo HAGO").italic = True

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
